# -*- coding: utf-8 -*-
"""
Robô de Auditoria de Preços — Itaueira
Cruza o faturamento com a Tabela de Preços oficial e aponta divergências.

Execução:
    pip install streamlit pandas numpy xlsxwriter
    streamlit run auditoria_precos.py
"""

import importlib
import io
import json
import re
import unicodedata
import zipfile
from datetime import date

import numpy as np
import pandas as pd
import streamlit as st

# =============================================================================
# CONFIGURAÇÃO DE NEGÓCIO
# =============================================================================

# De/Para oficial: tabela do faturamento -> tabela de preços.
# A lista é uma ordem de preferência. Uma mesma tabela do faturamento aponta
# para nomes diferentes conforme a linha de produto (melão usa "Merc Interno",
# pimentão/uva usam "FOB" e "CIF SP"). O robô escolhe o primeiro nome que
# realmente existe na tabela oficial para aquela variedade e marca, então não é
# preciso manter uma lista de quais variedades são de qual família.
DEPARA_OFICIAL = {
    "CFI Mercado Interno FOB - Sem Frete": ["4. Merc Int Sem Frete", "1. NO/NE FOB", "0. FOB"],
    "CFI Mercado Interno": ["1. Merc Interno", "2. SU/SE/CO CIF SP"],
    "CFI Mercado Interno NO e RJ": ["2. Merc Int (NOR, RJ, MT, MS)", "1. NO/NE FOB"],
    "CFI Mercado Interno SUL": ["3. Merc Int SUL", "2. SU/SE/CO CIF SP"],
    "CFI Mercado Interno Desconto 7%": ["6. DescFinanc 7%"],
    "CFI Mercado Interno Desconto 7% PR e RJ": ["7. DescFinanc 7% PR+RJ"],
    "CFI Mercado Interno Desconto 5%": ["5. DescFinanc 5%"],
    "CFI Mercado Interno GPA": ["9. GPA"],
    "CFI Mercado Interno Especial": ["1.1 Merc Interno Esp"],
}

# Tabelas do faturamento que não são venda a cliente — não entram na auditoria.
# Estabelecimentos ITR é filial de faturamento (transferência interna).
TABELAS_FORA_ESCOPO_PADRAO = ["Estabelecimentos ITR"]

# Só para agrupar o relatório por zona, não interfere no cruzamento
REGIOES_UF = {
    "NE": ["CE", "BA", "PE", "PB", "RN", "AL", "SE", "MA", "PI"],
    "NORTE": ["AM", "PA", "AP", "RR", "RO", "AC", "TO"],
    "SUL": ["RS", "SC", "PR"],
    "SUDESTE": ["SP", "MG", "ES", "RJ"],
    "CENTRO-OESTE": ["GO", "DF", "MT", "MS"],
}
UF_PARA_REGIAO = {uf: reg for reg, ufs in REGIOES_UF.items() for uf in ufs}

COLS_FAT = {
    "data": "emissaomovdate",
    "filial": "origemdestino.c",
    "tabela": "Tabela",
    "uf": "cliente.uf.c",
    "grupo": "cliente.grupo.c",
    "cliente": "cliente.c",
    "produto": "recurso.n",
    "variedade": "recurso.variedade.c",
    "marca": "recurso.classemarca.n",
    "tipo": "Tipo",
    "peso": "Peso Caixa",
    "qtd": "QTD caixa",
    "preco": "Preço Caixa",
}
COLS_TAB = {
    "ini": "vigencia_inicio",
    "fim": "vigencia_fim",
    "variedade_label": "VARIEDADE",
    "tipo_label": "TIPO",
    "peso": "PESO CX",
    "tabela": "TABELA",
    "preco": "Preço Final CX",
    "variedade": "Variedade.c",
    "marca": "marca.n",
    "tipomin": "Tipomin",
    "tipomax": "Tipomax",
}

# =============================================================================
# PARSING
# =============================================================================


def ler_csv(arquivo) -> pd.DataFrame:
    """Lê CSV tolerando separador, encoding e BOM variados."""
    bruto = arquivo.getvalue() if hasattr(arquivo, "getvalue") else arquivo.read()
    ultimo_erro = None
    for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
        for sep in (",", ";", "\t"):
            try:
                df = pd.read_csv(io.BytesIO(bruto), dtype=str, sep=sep, encoding=enc)
                if df.shape[1] > 1:
                    df.columns = [str(c).strip().lstrip("﻿") for c in df.columns]
                    return df
            except Exception as e:  # noqa: BLE001
                ultimo_erro = e
    raise ValueError(f"Não consegui ler o arquivo. Último erro: {ultimo_erro}")


def para_numero(valor):
    """Converte texto pt-BR em número. Trata 'R$ 1.234,56', '10 Kg', '4,5', '1.234'."""
    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return np.nan
    if isinstance(valor, (int, float, np.number)):
        return float(valor)

    s = str(valor).strip()
    if not s or s.lower() in {"nan", "none", "-", "--"}:
        return np.nan

    s = s.replace("\xa0", " ")
    s = re.sub(r"(?i)\bkgs?\b", "", s)        # "10 Kg" -> "10"
    s = re.sub(r"(?i)r\$", "", s)
    s = re.sub(r"[^\d,.\-]", "", s)
    if not s:
        return np.nan

    tem_virgula, tem_ponto = "," in s, "." in s
    if tem_virgula and tem_ponto:
        # o separador decimal é o que aparece por último
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") \
            else s.replace(",", "")
    elif tem_virgula:
        s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")                # 1.234.567 -> milhar
    elif tem_ponto:
        inteiro, _, decimal = s.partition(".")
        if len(decimal) == 3 and len(inteiro) <= 3 and inteiro.lstrip("-").isdigit():
            s = inteiro + decimal             # 1.234 -> milhar, não 1,234
    return pd.to_numeric(s, errors="coerce")


def para_data(serie: pd.Series) -> pd.Series:
    """Converte datas aceitando dd/mm/aaaa e ISO."""
    d = pd.to_datetime(serie, format="%d/%m/%Y", errors="coerce")
    falta = d.isna() & serie.notna()
    if falta.any():
        d.loc[falta] = pd.to_datetime(serie[falta], errors="coerce", dayfirst=True)
    return d.dt.normalize()


def normalizar(valor) -> str:
    """Chave de cruzamento: sem acento, sem espaço duplo, maiúscula."""
    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return ""
    s = unicodedata.normalize("NFKD", str(valor))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def chave_tabela(valor) -> str:
    """Normaliza o nome da tabela do faturamento.

    Absorve variações de digitação ('FI' por 'CFI'), pontuação e espaços,
    para o De/Para não quebrar por causa de um caractere.
    """
    s = normalizar(valor)
    s = re.sub(r"^C?FI\s+", "", s)            # 'CFI ' / 'FI ' viram prefixo opcional
    return re.sub(r"[^A-Z0-9%]+", "", s)


def nome_filial(valor) -> str:
    """'ITR CBR Itaueira Agropecuária Canto do Buriti' -> 'CBR — Canto do Buriti'"""
    s = str(valor).strip()
    m = re.match(r"^ITR\s+([A-Z]{2,4})\s+Itaueira\s+Agropecu[áa]ria\s+(.+)$", s, re.I)
    return f"{m.group(1).upper()} — {m.group(2).strip()}" if m else (s or "Não informada")


def extrair_codigos_variedade(codigo, rotulo) -> list:
    """'BAM/BAM/BLR' + '8. Pimentão BVM/BAM/BLR' -> ['BAM', 'BLR', 'BVM']"""
    achados = {p.strip().upper() for p in str(codigo).split("/") if p.strip()}
    achados |= set(re.findall(r"\b[A-Z]{2,4}\b", normalizar(rotulo)))
    achados = {c for c in achados if c.isalpha() and 2 <= len(c) <= 4}
    return sorted(achados) or [str(codigo).strip().upper()]


def conferir_colunas(df: pd.DataFrame, mapa: dict, nome: str):
    faltam = [c for c in mapa.values() if c not in df.columns]
    if faltam:
        raise KeyError(
            f"O arquivo de {nome} não tem as colunas: {', '.join(faltam)}.\n"
            f"Colunas encontradas: {', '.join(map(str, df.columns))}"
        )


# =============================================================================
# PREPARO DAS BASES
# =============================================================================


@st.cache_data(show_spinner=False)
def preparar_faturamento(bruto: bytes) -> pd.DataFrame:
    df = ler_csv(io.BytesIO(bruto))
    conferir_colunas(df, COLS_FAT, "faturamento")
    c = COLS_FAT
    out = pd.DataFrame(index=df.index)
    out["data"] = para_data(df[c["data"]])
    out["filial"] = df[c["filial"]].astype(str).str.strip().map(nome_filial)
    out["tabela_fat"] = df[c["tabela"]].astype(str).str.strip()
    out["tabela_key"] = out["tabela_fat"].map(chave_tabela)
    out["uf"] = df[c["uf"]].astype(str).str.strip().str.upper()
    out["regiao"] = out["uf"].map(UF_PARA_REGIAO).fillna("NÃO CLASSIFICADA")
    out["grupo"] = df[c["grupo"]].astype(str).str.strip()
    out["cliente"] = df[c["cliente"]].astype(str).str.strip()
    out["produto"] = df[c["produto"]].astype(str).str.strip()
    out["variedade"] = df[c["variedade"]].astype(str).str.strip().str.upper()
    out["marca"] = df[c["marca"]].astype(str).str.strip()
    out["marca_key"] = out["marca"].map(normalizar)
    out["tipo_txt"] = df[c["tipo"]].astype(str).str.strip()
    out["tipo"] = pd.to_numeric(out["tipo_txt"], errors="coerce")
    out["peso"] = df[c["peso"]].map(para_numero).round(4)
    out["qtd"] = df[c["qtd"]].map(para_numero)
    out["preco_fat"] = df[c["preco"]].map(para_numero)
    out["linha_csv"] = df.index + 2
    return out.reset_index(drop=True)


@st.cache_data(show_spinner=False)
def preparar_tabela(bruto: bytes) -> pd.DataFrame:
    df = ler_csv(io.BytesIO(bruto))
    conferir_colunas(df, COLS_TAB, "tabela de preços")
    c = COLS_TAB
    out = pd.DataFrame(index=df.index)
    out["vig_ini"] = para_data(df[c["ini"]])
    out["vig_fim"] = para_data(df[c["fim"]])
    out["tabela_of"] = df[c["tabela"]].astype(str).str.strip()
    out["variedade_raw"] = df[c["variedade"]].astype(str).str.strip()
    out["variedade_label"] = df[c["variedade_label"]].astype(str).str.strip()
    out["marca"] = df[c["marca"]].astype(str).str.strip()
    out["marca_key"] = out["marca"].map(normalizar)
    out["tipo_label"] = df[c["tipo_label"]].astype(str).str.strip()
    out["peso"] = df[c["peso"]].map(para_numero).round(4)
    out["preco_tab"] = df[c["preco"]].map(para_numero)
    out["tipo_min"] = df[c["tipomin"]].map(para_numero)
    out["tipo_max"] = df[c["tipomax"]].map(para_numero)

    # uma variedade composta (BVM/BAM/BLR) vira várias linhas, uma por código
    out["variedade"] = out.apply(
        lambda r: extrair_codigos_variedade(r["variedade_raw"], r["variedade_label"]),
        axis=1,
    )
    return out.explode("variedade").reset_index(drop=True)


# =============================================================================
# REGRA DE PRECIFICAÇÃO
# =============================================================================


def aplicar_depara(fat: pd.DataFrame, tab: pd.DataFrame, depara: dict,
                   fora_escopo: list) -> pd.DataFrame:
    """Define, para cada linha faturada, qual linha da tabela oficial é a correta.

    Quando o De/Para oferece mais de um destino (caso de FOB, que se chama
    "4. Merc Int Sem Frete" no melão e "1. NO/NE FOB" no pimentão), vence o
    primeiro destino que existe na tabela oficial para aquela variedade e marca.
    """
    fat = fat.copy()
    depara_key = {chave_tabela(k): list(v) for k, v in depara.items()}
    fora_key = {chave_tabela(x) for x in fora_escopo}

    # quais tabelas existem para cada variedade + marca
    disponiveis = (
        tab.dropna(subset=["preco_tab"])
        .groupby(["variedade", "marca_key"])["tabela_of"].agg(set).to_dict()
    )

    def resolver(r):
        if r["tabela_key"] in fora_key:
            return None, "Fora de escopo (não é venda a cliente)"
        candidatos = depara_key.get(r["tabela_key"])
        if not candidatos:
            return None, f"Tabela '{r['tabela_fat']}' sem De/Para configurado"
        existentes = disponiveis.get((r["variedade"], r["marca_key"]), set())
        for destino in candidatos:
            if destino in existentes:
                return destino, "De/Para oficial"
        # nenhum candidato existe para este produto: usa o primeiro para que a
        # linha apareça como "sem preço na tabela" em vez de sumir do relatório
        return candidatos[0], "De/Para oficial (produto sem preço cadastrado)"

    res = fat.apply(resolver, axis=1, result_type="expand")
    fat["tabela_alvo"] = res[0]
    fat["origem_regra"] = res[1]
    return fat


def auditar(fat: pd.DataFrame, tab: pd.DataFrame,
            tol_reais: float, tol_pct: float) -> pd.DataFrame:
    """LEFT join: nenhuma linha some — cada uma recebe um status."""
    auditavel = fat[fat["tabela_alvo"].notna()].copy()
    auditavel["_id"] = np.arange(len(auditavel))

    m = auditavel.merge(
        tab,
        left_on=["variedade", "marca_key", "tabela_alvo", "peso"],
        right_on=["variedade", "marca_key", "tabela_of", "peso"],
        how="left", suffixes=("", "_tab"),
    )

    tem_candidato = m["preco_tab"].notna()
    na_vigencia = (m["data"] >= m["vig_ini"]) & (m["data"] <= m["vig_fim"])
    # linhas sem faixa de tipo (pimentão: 'Vários', '250g') cruzam só por peso
    sem_faixa = m["tipo_min"].isna() | m["tipo_max"].isna()
    na_faixa = sem_faixa | ((m["tipo"] >= m["tipo_min"]) & (m["tipo"] <= m["tipo_max"]))
    valido = tem_candidato & na_vigencia & na_faixa

    # 1 preço por linha faturada; se nada validou, mantém a linha sem preço
    bons = m[valido].drop_duplicates(subset="_id", keep="first")
    orfas = m[~m["_id"].isin(bons["_id"])].drop_duplicates(subset="_id", keep="first")
    for col in ["preco_tab", "vig_ini", "vig_fim", "tipo_label", "tipo_min", "tipo_max"]:
        orfas[col] = np.nan

    r = pd.concat([bons, orfas], ignore_index=True).sort_values("_id")
    r["preco_fat"] = r["preco_fat"].round(2)
    r["preco_tab"] = r["preco_tab"].round(2)
    r["diferenca"] = (r["preco_fat"] - r["preco_tab"]).round(2)
    r["dif_pct"] = np.where(
        r["preco_tab"].fillna(0) != 0, r["diferenca"] / r["preco_tab"] * 100, np.nan
    ).round(2)
    r["impacto_rs"] = (r["diferenca"] * r["qtd"]).round(2)

    dentro = (r["diferenca"].abs() <= tol_reais) | (r["dif_pct"].abs() <= tol_pct)
    r["status"] = np.select(
        [r["preco_tab"].isna(), dentro, r["diferenca"] > 0, r["diferenca"] < 0],
        ["SEM PREÇO NA TABELA", "OK", "FATURADO A MAIOR", "FATURADO A MENOR"],
        default="OK",
    )

    nao_auditadas = fat[fat["tabela_alvo"].isna()].copy()
    for col in ["preco_tab", "diferenca", "dif_pct", "impacto_rs"]:
        nao_auditadas[col] = np.nan
    nao_auditadas["status"] = "NÃO AUDITADO"
    return pd.concat([r, nao_auditadas], ignore_index=True)


def identificar_origem_do_preco(res: pd.DataFrame, tab: pd.DataFrame,
                                tol_reais: float) -> pd.Series:
    """Para cada linha, diz a qual coluna da tabela o preço FATURADO corresponde.

    Separa 'aplicaram a tabela errada' (o preço bate com outra coluna vigente)
    de 'preço fora de tabela' (não bate com nenhuma) — o primeiro é erro de
    rota/cadastro, o segundo é desconto comercial a justificar.
    """
    cand = res.reset_index(drop=True).assign(_lin=lambda d: d.index).merge(
        tab, on=["variedade", "marca_key", "peso"], how="left", suffixes=("", "_c"),
    )
    cand = cand[(cand["data"] >= cand["vig_ini_c"]) & (cand["data"] <= cand["vig_fim_c"])]
    sem_faixa = cand["tipo_min_c"].isna() | cand["tipo_max_c"].isna()
    cand = cand[sem_faixa | ((cand["tipo"] >= cand["tipo_min_c"])
                             & (cand["tipo"] <= cand["tipo_max_c"]))]
    bate = cand[(cand["preco_fat"] - cand["preco_tab_c"]).abs() <= tol_reais]
    achado = bate.groupby("_lin")["tabela_of_c"].apply(lambda s: " / ".join(sorted(set(s))))
    return achado.reindex(range(len(res))).fillna("")


def classificar_motivo(res: pd.DataFrame) -> pd.Series:
    """Traduz o cruzamento em uma causa provável, para direcionar a cobrança."""
    def motivo(r):
        if r["status"] == "NÃO AUDITADO":
            return r["origem_regra"]
        if r["status"] == "SEM PREÇO NA TABELA":
            return "Combinação não cadastrada na tabela vigente"
        if r["status"] == "OK":
            return "Conforme"
        if r["confere_com"]:
            return f"Tabela aplicada divergente — o preço é o de {r['confere_com']}"
        return "Preço fora de qualquer tabela vigente (desconto comercial?)"
    return res.apply(motivo, axis=1)


# =============================================================================
# SAÍDAS
# =============================================================================

COLS_SAIDA = [
    "data", "filial", "cliente", "grupo", "uf", "regiao", "produto", "variedade", "marca",
    "tipo_txt", "peso", "qtd", "tabela_fat", "tabela_alvo", "tipo_label",
    "preco_fat", "preco_tab", "diferenca", "dif_pct", "impacto_rs", "status",
    "motivo", "linha_csv",
]
ROTULOS = {
    "data": "Emissão", "filial": "Filial de faturamento", "cliente": "Cliente", "grupo": "Grupo", "uf": "UF",
    "regiao": "Região", "produto": "Produto", "variedade": "Variedade",
    "marca": "Marca", "tipo_txt": "Tipo", "peso": "Peso Cx", "qtd": "Qtd Cx",
    "tabela_fat": "Tabela (faturamento)", "tabela_alvo": "Tabela oficial aplicada",
    "tipo_label": "Faixa da tabela", "preco_fat": "Preço faturado",
    "preco_tab": "Preço tabela", "diferenca": "Diferença R$",
    "dif_pct": "Diferença %", "impacto_rs": "Impacto R$", "status": "Status",
    "motivo": "Causa provável", "linha_csv": "Linha do CSV",
}


def formatar(df: pd.DataFrame) -> pd.DataFrame:
    out = df.reindex(columns=[c for c in COLS_SAIDA if c in df.columns]).copy()
    if "data" in out:
        out["data"] = pd.to_datetime(out["data"]).dt.strftime("%d/%m/%Y")
    return out.rename(columns=ROTULOS)


def brl(valor) -> str:
    """Formata no padrão pt-BR: 1.234.567,89"""
    if valor is None or pd.isna(valor):
        return "0,00"
    return f"{valor:,.2f}".replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def motor_excel():
    """Devolve o primeiro motor de Excel disponível. Streamlit Cloud costuma ter
    apenas openpyxl; o app não pode quebrar por causa disso."""
    for nome in ("xlsxwriter", "openpyxl"):
        try:
            importlib.import_module(nome)
            return nome
        except ImportError:
            continue
    return None


def gerar_excel(abas: list) -> bytes:
    motor = motor_excel()
    if motor is None:
        raise ModuleNotFoundError(
            "Nenhum motor de Excel instalado. Adicione 'xlsxwriter' ou 'openpyxl' "
            "ao requirements.txt do app."
        )
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine=motor) as xl:
        for nome, dados in abas:
            if dados is None or dados.empty:
                dados = pd.DataFrame({"Sem registros": []})
            aba = re.sub(r"[\[\]:*?/\\]", "-", str(nome))[:31]
            larguras = [
                min(max(len(str(c)) + 2,
                        (10 if pd.isna(dados[c].astype(str).str.len().max())
                         else int(dados[c].astype(str).str.len().max())) + 2), 45)
                for c in dados.columns
            ]
            if motor == "xlsxwriter":
                dados.to_excel(xl, sheet_name=aba, index=False, startrow=1, header=False)
                livro, ws = xl.book, xl.sheets[aba]
                cab = livro.add_format({"bold": True, "bg_color": "#1F4E79",
                                        "font_color": "white", "border": 1,
                                        "align": "center", "valign": "vcenter"})
                moeda = livro.add_format({"num_format": "#,##0.00"})
                for i, col in enumerate(dados.columns):
                    ws.write(0, i, str(col), cab)
                    ws.set_column(i, i, larguras[i],
                                  moeda if dados[col].dtype.kind in "fc" else None)
                ws.freeze_panes(1, 0)
                if len(dados):
                    ws.autofilter(0, 0, len(dados), len(dados.columns) - 1)
            else:
                from openpyxl.styles import Alignment, Font, PatternFill
                from openpyxl.utils import get_column_letter
                dados.to_excel(xl, sheet_name=aba, index=False)
                ws = xl.sheets[aba]
                preenche = PatternFill("solid", fgColor="1F4E79")
                for i, col in enumerate(dados.columns, start=1):
                    celula = ws.cell(row=1, column=i)
                    celula.font = Font(bold=True, color="FFFFFF")
                    celula.fill = preenche
                    celula.alignment = Alignment(horizontal="center", vertical="center")
                    ws.column_dimensions[get_column_letter(i)].width = larguras[i - 1]
                    if dados[col].dtype.kind in "fc":
                        for linha in range(2, len(dados) + 2):
                            ws.cell(row=linha, column=i).number_format = "#,##0.00"
                ws.freeze_panes = "A2"
                if len(dados):
                    ws.auto_filter.ref = (
                        f"A1:{get_column_letter(len(dados.columns))}{len(dados) + 1}")
    return buf.getvalue()


def gerar_csvs_zip(abas: list) -> bytes:
    """Plano B quando nenhum motor de Excel está disponível."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in abas:
            if dados is None or dados.empty:
                continue
            arquivo = re.sub(r"[^\w\- ]+", "", str(nome)).strip().replace(" ", "_")
            z.writestr(f"{arquivo}.csv",
                       dados.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"))
    return buf.getvalue()


# --- E-mail -----------------------------------------------------------------

def _tabela_texto(df: pd.DataFrame, alinhar_dir: list) -> str:
    """Tabela em largura fixa, legível em e-mail de texto puro."""
    if df.empty:
        return "  (sem registros)"
    larg = {c: max(len(str(c)), *(len(str(v)) for v in df[c])) for c in df.columns}
    def linha(vals):
        return "  " + "  ".join(
            str(v).rjust(larg[c]) if c in alinhar_dir else str(v).ljust(larg[c])
            for c, v in zip(df.columns, vals))
    sep = "  " + "  ".join("-" * larg[c] for c in df.columns)
    return "\n".join([linha(df.columns), sep] + [linha(r) for r in df.itertuples(index=False)])


def resumo_por_cliente(div: pd.DataFrame, limite: int) -> pd.DataFrame:
    g = (div.groupby(["cliente", "uf"], as_index=False)
         .agg(lanc=("status", "size"), caixas=("qtd", "sum"), impacto=("impacto_rs", "sum")))
    g = g.reindex(g["impacto"].abs().sort_values(ascending=False).index)
    corpo = g.head(limite)
    saida = pd.DataFrame({
        "Cliente": corpo["cliente"], "UF": corpo["uf"],
        "Lanç.": corpo["lanc"].astype(int),
        "Caixas": corpo["caixas"].map(lambda v: brl(v).replace(",00", "")),
        "Impacto R$": corpo["impacto"].map(brl),
    })
    if len(g) > limite:
        resto = g.tail(len(g) - limite)
        saida.loc[len(saida)] = [f"+ outros {len(resto)} cliente(s)", "",
                                 int(resto["lanc"].sum()),
                                 brl(resto["caixas"].sum()).replace(",00", ""),
                                 brl(resto["impacto"].sum())]
    saida.loc[len(saida)] = ["TOTAL", "", int(g["lanc"].sum()),
                             brl(g["caixas"].sum()).replace(",00", ""),
                             brl(g["impacto"].sum())]
    return saida


def blocos_data_filial(div: pd.DataFrame, itens_por_bloco: int) -> list:
    """Divergências agrupadas por data de emissão e filial de faturamento."""
    blocos = []
    for (dia, filial), g in div.groupby([div["data"].dt.date, "filial"], sort=True):
        g = g.reindex(g["impacto_rs"].abs().sort_values(ascending=False).index)
        itens = [
            {"Cliente": f"{r['cliente']} ({r['uf']})",
             "Produto": (f"{r['variedade']} {r['marca']} "
                         f"{'T' + r['tipo_txt'] if str(r['tipo_txt']).isdigit() else r['tipo_txt']}"
                         f" cx{r['peso']:g}kg"),
             "Faturado": brl(r["preco_fat"]), "Tabela": brl(r["preco_tab"]),
             "Dif.": brl(r["diferenca"]), "Qtd": brl(r["qtd"]).replace(",00", ""),
             "Impacto R$": brl(r["impacto_rs"]), "Tabela aplicada": r["tabela_alvo"]}
            for _, r in g.head(itens_por_bloco).iterrows()
        ]
        blocos.append({
            "data": dia.strftime("%d/%m/%Y"), "filial": filial,
            "lanc": len(g), "impacto": g["impacto_rs"].sum(),
            "ocultos": max(len(g) - itens_por_bloco, 0),
            "itens": pd.DataFrame(itens),
        })
    return sorted(blocos, key=lambda b: (b["data"], b["filial"]))


def gerar_email(div: pd.DataFrame, periodo: str, destinatarios: str,
                total_conferido: int, top_clientes: int = 12,
                itens_por_bloco: int = 5) -> str:
    if div.empty:
        return "Nenhuma divergência encontrada no período."

    menor = div[div["status"] == "FATURADO A MENOR"]
    maior = div[div["status"] == "FATURADO A MAIOR"]
    troca = div[div["confere_com"] != ""]
    fora = div[div["confere_com"] == ""]

    partes = [
        "Prezados,",
        "",
        "Pedimos a sua verificação para os preços com as diferenças apresentadas nos "
        "quadros abaixo. Solicitamos os seus comentários.",
        "",
        "=" * 78,
        "RESUMO",
        "=" * 78,
        f"  Período auditado ................ {periodo}",
        f"  Lançamentos conferidos .......... {total_conferido}",
        f"  Lançamentos com divergência ..... {len(div)}",
        f"  Faturados a MENOR que a tabela .. {len(menor)}  |  "
        f"R$ {brl(menor['impacto_rs'].sum())}",
        f"  Faturados a MAIOR que a tabela .. {len(maior)}  |  "
        f"R$ {brl(maior['impacto_rs'].sum())}",
        f"  IMPACTO LÍQUIDO ................. R$ {brl(div['impacto_rs'].sum())}",
        "",
        f"  Causa provável: {len(troca)} lançamento(s) com preço de OUTRA tabela vigente "
        "(rota/cadastro)",
        f"                  {len(fora)} lançamento(s) com preço fora de qualquer tabela "
        "(desconto comercial?)",
        "",
        "=" * 78,
        "1) IMPACTO POR CLIENTE",
        "=" * 78,
        _tabela_texto(resumo_por_cliente(div, top_clientes),
                      ["Lanç.", "Caixas", "Impacto R$"]),
        "",
        "=" * 78,
        "2) DISTORÇÕES POR DATA DE EMISSÃO E FILIAL DE FATURAMENTO",
        "=" * 78,
    ]

    for b in blocos_data_filial(div, itens_por_bloco):
        partes += [
            "",
            f"► {b['data']}  |  {b['filial']}",
            f"  {b['lanc']} lançamento(s)  |  impacto R$ {brl(b['impacto'])}",
            _tabela_texto(b["itens"], ["Faturado", "Tabela", "Dif.", "Qtd", "Impacto R$"]),
        ]
        if b["ocultos"]:
            partes.append(f"  ... e mais {b['ocultos']} lançamento(s) desta data/filial "
                          "no anexo.")

    partes += [
        "",
        "=" * 78,
        "",
        "O detalhamento completo, linha a linha, está na planilha em anexo.",
        "Favor informar se há desconto ou condição comercial aprovada para estes casos.",
        "",
        "Atenciosamente,",
        "Auditoria de Preços",
        "",
        f"Para: {destinatarios}",
    ]
    return "\n".join(partes)


def gerar_email_html(div: pd.DataFrame, periodo: str, destinatarios: str,
                     total_conferido: int, top_clientes: int = 12,
                     itens_por_bloco: int = 5) -> str:
    """Mesma estrutura, em HTML — cola no Gmail mantendo as tabelas."""
    if div.empty:
        return "<p>Nenhuma divergência encontrada no período.</p>"

    menor = div[div["status"] == "FATURADO A MENOR"]
    maior = div[div["status"] == "FATURADO A MAIOR"]
    troca = div[div["confere_com"] != ""]
    fora = div[div["confere_com"] == ""]

    est_tab = ("border-collapse:collapse;font:13px Arial,sans-serif;"
               "margin:6px 0 16px 0;width:100%")
    est_th = ("background:#1F4E79;color:#fff;padding:6px 9px;text-align:left;"
              "border:1px solid #cfd8e3;white-space:nowrap")
    est_td = "padding:5px 9px;border:1px solid #cfd8e3;vertical-align:top"

    def tabela(df, alinhar_dir, negrito_ultima=False):
        if df.empty:
            return "<p style='color:#666'>(sem registros)</p>"
        cab = "".join(f"<th style='{est_th}'>{c}</th>" for c in df.columns)
        linhas = []
        for i, r in enumerate(df.itertuples(index=False)):
            ult = negrito_ultima and i == len(df) - 1
            fundo = "background:#eef3f9;font-weight:bold" if ult else (
                "background:#fafbfd" if i % 2 else "")
            tds = "".join(
                f"<td style='{est_td};{fundo};"
                f"{'text-align:right' if c in alinhar_dir else ''}'>{v}</td>"
                for c, v in zip(df.columns, r))
            linhas.append(f"<tr>{tds}</tr>")
        return (f"<table style='{est_tab}'><thead><tr>{cab}</tr></thead>"
                f"<tbody>{''.join(linhas)}</tbody></table>")

    cor = "#c0392b" if div["impacto_rs"].sum() < 0 else "#1e8449"
    html = [
        "<div style=\"font:14px Arial,sans-serif;color:#1a1a1a;max-width:1000px\">",
        "<p>Prezados,</p>",
        "<p>Pedimos a sua verificação para os preços com as diferenças apresentadas "
        "nos quadros abaixo. Solicitamos os seus comentários.</p>",
        "<h3 style='margin:18px 0 4px;color:#1F4E79'>Resumo</h3>",
        tabela(pd.DataFrame({
            "Indicador": ["Período auditado", "Lançamentos conferidos",
                          "Lançamentos com divergência", "Faturados a MENOR que a tabela",
                          "Faturados a MAIOR que a tabela", "Impacto líquido"],
            "Valor": [periodo, f"{total_conferido}", f"{len(div)}",
                      f"{len(menor)} &nbsp;|&nbsp; R$ {brl(menor['impacto_rs'].sum())}",
                      f"{len(maior)} &nbsp;|&nbsp; R$ {brl(maior['impacto_rs'].sum())}",
                      f"<b style='color:{cor}'>R$ {brl(div['impacto_rs'].sum())}</b>"],
        }), []),
        f"<p style='font-size:13px;color:#444'><b>Causa provável:</b> {len(troca)} "
        f"lançamento(s) com preço de <b>outra tabela vigente</b> (rota/cadastro) e "
        f"{len(fora)} com preço <b>fora de qualquer tabela</b> (desconto comercial?).</p>",
        "<h3 style='margin:22px 0 4px;color:#1F4E79'>1) Impacto por cliente</h3>",
        tabela(resumo_por_cliente(div, top_clientes),
               ["Lanç.", "Caixas", "Impacto R$"], negrito_ultima=True),
        "<h3 style='margin:22px 0 4px;color:#1F4E79'>2) Distorções por data de emissão "
        "e filial de faturamento</h3>",
    ]
    for b in blocos_data_filial(div, itens_por_bloco):
        html.append(
            f"<div style='margin-top:14px;padding:6px 10px;background:#eef3f9;"
            f"border-left:4px solid #1F4E79;font-size:13px'>"
            f"<b>{b['data']}</b> &nbsp;·&nbsp; {b['filial']} &nbsp;·&nbsp; "
            f"{b['lanc']} lançamento(s) &nbsp;·&nbsp; impacto "
            f"<b>R$ {brl(b['impacto'])}</b></div>")
        html.append(tabela(b["itens"], ["Faturado", "Tabela", "Dif.", "Qtd", "Impacto R$"]))
        if b["ocultos"]:
            html.append(f"<p style='font-size:12px;color:#666;margin-top:-10px'>"
                        f"... e mais {b['ocultos']} lançamento(s) desta data/filial "
                        f"no anexo.</p>")
    html += [
        "<p style='margin-top:20px'>O detalhamento completo, linha a linha, está na "
        "planilha em anexo.<br>Favor informar se há desconto ou condição comercial "
        "aprovada para estes casos.</p>",
        "<p>Atenciosamente,<br><b>Auditoria de Preços</b></p>",
        f"<p style='font-size:12px;color:#666'>Para: {destinatarios}</p>",
        "</div>",
    ]
    return "".join(html)


# =============================================================================
# INTERFACE
# =============================================================================

st.set_page_config(page_title="Auditoria de Preços", layout="wide", page_icon="🤖")
st.title("🤖 Robô de Auditoria de Preços")
st.caption("Cruza o faturamento com a Tabela de Preços oficial e aponta as divergências.")

with st.sidebar:
    st.header("⚙️ Parâmetros")
    tol_reais = st.number_input("Tolerância em R$ por caixa", 0.0, 50.0, 0.01, 0.01,
                                help="Diferenças até este valor são tratadas como arredondamento.")
    tol_pct = st.number_input("Tolerância em %", 0.0, 20.0, 0.0, 0.1)
    st.divider()
    fora_escopo = st.multiselect(
        "Tabelas fora do escopo", options=TABELAS_FORA_ESCOPO_PADRAO + ["Clientes Diversos"],
        default=TABELAS_FORA_ESCOPO_PADRAO,
        help="Estabelecimentos ITR é filial de faturamento — transferência interna, não venda.",
    )
    st.divider()
    config_json = st.file_uploader("Restaurar De/Para salvo (JSON)", type=["json"])
    destinatarios = st.text_input(
        "Destinatários do e-mail",
        "com.inteligencia.de.mercado@itaueira.com; dir.comercial@itaueira.com",
    )

col1, col2 = st.columns(2)
with col1:
    arq_fat = st.file_uploader("📥 CSV do Faturamento", type=["csv"])
with col2:
    arq_tab = st.file_uploader("📥 CSV da Tabela de Preços", type=["csv"])

if not (arq_fat and arq_tab):
    st.info("Carregue os dois arquivos para iniciar a auditoria.")
    st.stop()

try:
    fat = preparar_faturamento(arq_fat.getvalue())
    tab = preparar_tabela(arq_tab.getvalue())
except Exception as erro:  # noqa: BLE001
    st.error(f"Não consegui preparar os arquivos.\n\n**{type(erro).__name__}:** {erro}")
    st.stop()

# --- De/Para editável -------------------------------------------------------
depara_inicial = DEPARA_OFICIAL
if config_json is not None:
    try:
        depara_inicial = json.loads(config_json.getvalue().decode("utf-8"))["depara"]
        st.sidebar.success("De/Para restaurado do arquivo.")
    except Exception as erro:  # noqa: BLE001
        st.sidebar.error(f"JSON inválido: {erro}")

with st.expander("🔗 De/Para oficial — clique para conferir ou ajustar", expanded=False):
    st.markdown(
        "Uma mesma tabela do faturamento pode ter nomes diferentes na tabela de preços "
        "conforme a linha de produto (melão usa *Merc Interno*, pimentão e uva usam *FOB* "
        "e *CIF SP*). Use a **ordem** para dizer a preferência: o robô aplica o primeiro "
        "destino que existe na tabela oficial para aquela variedade e marca."
    )
    opcoes_tabela = sorted(tab["tabela_of"].dropna().unique().tolist())
    base = pd.DataFrame(
        [{"Tabela no faturamento": k, "Ordem": i + 1, "Tabela de preço": v}
         for k, destinos in depara_inicial.items() for i, v in enumerate(destinos)]
    )
    editado = st.data_editor(
        base, hide_index=True, use_container_width=True, num_rows="dynamic", key="ed_depara",
        column_config={
            "Tabela no faturamento": st.column_config.TextColumn(required=True),
            "Ordem": st.column_config.NumberColumn(min_value=1, step=1, required=True),
            "Tabela de preço": st.column_config.SelectboxColumn(
                options=opcoes_tabela, required=True),
        },
    )
    depara = {
        chave: grupo.sort_values("Ordem")["Tabela de preço"].tolist()
        for chave, grupo in editado.dropna(subset=["Tabela no faturamento"])
        .groupby("Tabela no faturamento")
    }

    faltando = sorted(
        set(fat["tabela_fat"]) - {k for k in depara} - set(fora_escopo)
        - {t for t in fat["tabela_fat"] if chave_tabela(t) in {chave_tabela(k) for k in depara}}
    )
    if faltando:
        st.warning("Sem De/Para (vão para *Não auditado*): " + ", ".join(faltando))

    st.download_button(
        "💾 Salvar este De/Para (JSON)",
        json.dumps({"depara": depara, "fora_escopo": fora_escopo},
                   ensure_ascii=False, indent=2).encode("utf-8"),
        file_name="config_auditoria_precos.json", mime="application/json",
    )

# --- Execução ---------------------------------------------------------------
with st.spinner("Cruzando faturamento x tabela oficial..."):
    fat_regra = aplicar_depara(fat, tab, depara, fora_escopo)
    res = auditar(fat_regra, tab, tol_reais, tol_pct).reset_index(drop=True)
    res["confere_com"] = identificar_origem_do_preco(res, tab, max(tol_reais, 0.01))
    res["motivo"] = classificar_motivo(res)

divergencias = res[res["status"].isin(["FATURADO A MAIOR", "FATURADO A MENOR"])].copy()
sem_preco = res[res["status"] == "SEM PREÇO NA TABELA"].copy()
nao_auditado = res[res["status"] == "NÃO AUDITADO"].copy()
conferidas = res[res["status"] != "NÃO AUDITADO"]

datas = fat["data"].dropna()
periodo = f"{datas.min():%d/%m/%Y} a {datas.max():%d/%m/%Y}" if len(datas) else "n/d"

st.divider()
st.subheader(f"Resultado — período {periodo}")
k = st.columns(6)
k[0].metric("Linhas no faturamento", f"{len(res):,}".replace(",", "."))
k[1].metric("Conferidas", f"{len(conferidas):,}".replace(",", "."))
k[2].metric("✅ Conformes", f"{(res['status'] == 'OK').sum():,}".replace(",", "."))
k[3].metric("⚠️ Divergentes", f"{len(divergencias):,}".replace(",", "."),
            delta=f"{len(divergencias) / max(len(conferidas), 1) * 100:.1f}% do conferido",
            delta_color="inverse")
k[4].metric("❓ Sem preço na tabela", f"{len(sem_preco):,}".replace(",", "."))
k[5].metric("💰 Impacto líquido", f"R$ {brl(divergencias['impacto_rs'].sum(skipna=True))}")

if divergencias.empty and sem_preco.empty:
    st.success("✅ Nenhuma divergência encontrada no período.")

aba1, aba2, aba3, aba4, aba5 = st.tabs(
    ["⚠️ Divergências", "📊 Resumos", "❓ Sem preço", "🚫 Não auditado", "🛠️ Diagnóstico"]
)

with aba1:
    if divergencias.empty:
        st.success("Nenhuma divergência acima da tolerância.")
    else:
        c1, c2 = st.columns([1, 2])
        f_status = c1.multiselect("Status", ["FATURADO A MENOR", "FATURADO A MAIOR"],
                                  default=["FATURADO A MENOR", "FATURADO A MAIOR"])
        causas = sorted(divergencias["motivo"].unique())
        f_causa = c2.multiselect("Causa provável", causas, default=causas)
        vis = divergencias[divergencias["status"].isin(f_status)
                           & divergencias["motivo"].isin(f_causa)]
        vis = vis.reindex(vis["impacto_rs"].abs().sort_values(ascending=False).index)
        st.caption(f"{len(vis)} lançamento(s) · impacto R$ {brl(vis['impacto_rs'].sum())}")
        st.dataframe(formatar(vis), use_container_width=True, hide_index=True, height=430)

resumo_cliente = resumo_causa = resumo_regiao = pd.DataFrame()
resumo_filial = resumo_data = pd.DataFrame()
with aba2:
    if divergencias.empty:
        st.info("Sem divergências para resumir.")
    else:
        def resumir(chaves, nomes):
            return (divergencias.groupby(chaves, as_index=False)
                    .agg(Lançamentos=("status", "size"), Caixas=("qtd", "sum"),
                         Impacto_RS=("impacto_rs", "sum"))
                    .sort_values("Impacto_RS", key=abs, ascending=False)
                    .rename(columns={**nomes, "Impacto_RS": "Impacto R$"}))

        resumo_cliente = resumir(["cliente", "uf"], {"cliente": "Cliente", "uf": "UF"})
        resumo_causa = resumir(["motivo", "status"], {"motivo": "Causa provável",
                                                      "status": "Status"})
        resumo_regiao = resumir(["regiao"], {"regiao": "Região"})
        resumo_filial = resumir(["filial"], {"filial": "Filial de faturamento"})
        resumo_data = (divergencias.assign(Emissão=divergencias["data"].dt.strftime("%d/%m/%Y"))
                       .groupby(["Emissão", "filial"], as_index=False)
                       .agg(Lançamentos=("status", "size"), Caixas=("qtd", "sum"),
                            Impacto_RS=("impacto_rs", "sum"))
                       .rename(columns={"filial": "Filial de faturamento",
                                        "Impacto_RS": "Impacto R$"})
                       .sort_values(["Emissão", "Filial de faturamento"]))

        e1, e2 = st.columns(2)
        with e1:
            st.markdown("**Impacto por cliente**")
            st.dataframe(resumo_cliente, use_container_width=True, hide_index=True, height=300)
            st.markdown("**Por filial de faturamento**")
            st.dataframe(resumo_filial, use_container_width=True, hide_index=True)
            st.markdown("**Por região**")
            st.dataframe(resumo_regiao, use_container_width=True, hide_index=True)
        with e2:
            st.markdown("**Por data de emissão e filial**")
            st.dataframe(resumo_data, use_container_width=True, hide_index=True, height=300)
            st.markdown("**Por causa provável**")
            st.dataframe(resumo_causa, use_container_width=True, hide_index=True)
            st.markdown("**Por variedade e marca**")
            st.dataframe(resumir(["variedade", "marca"],
                                 {"variedade": "Variedade", "marca": "Marca"}),
                         use_container_width=True, hide_index=True)

with aba3:
    st.caption("Combinações faturadas que não existem na tabela oficial vigente. "
               "Não é divergência de preço — é lacuna de cadastro.")
    if sem_preco.empty:
        st.success("Todas as combinações faturadas têm preço na tabela.")
    else:
        st.dataframe(
            sem_preco.groupby(["variedade", "marca", "peso", "tipo_txt", "tabela_alvo"],
                              as_index=False)
            .agg(Lançamentos=("status", "size"), Caixas=("qtd", "sum"))
            .rename(columns={"variedade": "Variedade", "marca": "Marca", "peso": "Peso Cx",
                             "tipo_txt": "Tipo", "tabela_alvo": "Tabela procurada"}),
            use_container_width=True, hide_index=True,
        )

with aba4:
    st.caption("Linhas deliberadamente fora da auditoria, ou sem De/Para configurado.")
    if nao_auditado.empty:
        st.info("Nenhuma linha fora do escopo.")
    else:
        st.dataframe(
            nao_auditado.groupby(["tabela_fat", "origem_regra"], as_index=False)
            .agg(Lançamentos=("status", "size"), Caixas=("qtd", "sum"))
            .rename(columns={"tabela_fat": "Tabela (faturamento)", "origem_regra": "Motivo"}),
            use_container_width=True, hide_index=True,
        )

with aba5:
    st.markdown("**Cobertura do cruzamento**")
    st.dataframe(res["status"].value_counts().rename_axis("Status").reset_index(name="Linhas"),
                 use_container_width=True, hide_index=True)
    d1, d2 = st.columns(2)
    with d1:
        st.markdown("**De/Para efetivamente aplicado**")
        st.dataframe(
            res[res["tabela_alvo"].notna()]
            .groupby(["tabela_fat", "tabela_alvo"], as_index=False).size()
            .rename(columns={"tabela_fat": "Tabela (faturamento)",
                             "tabela_alvo": "Tabela de preço", "size": "Linhas"}),
            use_container_width=True, hide_index=True, height=300,
        )
        st.markdown("**Vigências na tabela oficial**")
        st.dataframe(
            tab.groupby(["vig_ini", "vig_fim"], as_index=False).size()
            .assign(vig_ini=lambda d: d["vig_ini"].dt.strftime("%d/%m/%Y"),
                    vig_fim=lambda d: d["vig_fim"].dt.strftime("%d/%m/%Y"))
            .rename(columns={"vig_ini": "Início", "vig_fim": "Fim", "size": "Linhas"}),
            use_container_width=True, hide_index=True,
        )
        sem_vig = fat[~fat["data"].apply(
            lambda d: ((tab["vig_ini"] <= d) & (tab["vig_fim"] >= d)).any())]
        if len(sem_vig):
            st.warning(f"{len(sem_vig)} linha(s) faturadas fora de qualquer vigência: "
                       + ", ".join(sorted(sem_vig["data"].dt.strftime("%d/%m/%Y").unique())))
    with d2:
        st.markdown("**Valores que não viraram número**")
        st.dataframe(
            pd.DataFrame({
                "Campo": ["Preço faturado", "Peso caixa", "Qtd caixas", "Data emissão",
                          "Preço tabela", "Vigência início"],
                "Nulos": [fat["preco_fat"].isna().sum(), fat["peso"].isna().sum(),
                          fat["qtd"].isna().sum(), fat["data"].isna().sum(),
                          tab["preco_tab"].isna().sum(), tab["vig_ini"].isna().sum()],
            }), use_container_width=True, hide_index=True,
        )
        st.caption("Tipos não numéricos no faturamento ('Vários', '250g') são cruzados "
                   "por peso da caixa, pois a tabela não define faixa de tipo para eles.")

# --- Entregáveis ------------------------------------------------------------
st.divider()
st.subheader("📤 Entregáveis")

c1, c2 = st.columns(2)
top_clientes = c1.slider("Clientes no quadro de impacto", 5, 40, 12,
                         help="Os demais entram agregados como '+ outros N clientes'.")
itens_bloco = c2.slider("Itens detalhados por data/filial", 1, 20, 5,
                        help="O restante fica na planilha em anexo.")

texto_email = gerar_email(divergencias, periodo, destinatarios, len(conferidas),
                          top_clientes, itens_bloco)
html_email = gerar_email_html(divergencias, periodo, destinatarios, len(conferidas),
                              top_clientes, itens_bloco)

e_html, e_txt = st.tabs(["✉️ E-mail formatado (copiar e colar no Gmail)", "📝 Texto puro"])
with e_html:
    st.caption("Selecione tudo abaixo, copie (Ctrl+C) e cole no Gmail — as tabelas "
               "vão junto.")
    st.markdown(html_email, unsafe_allow_html=True)
with e_txt:
    st.text_area("Versão em texto", texto_email, height=420, label_visibility="collapsed")

abas_saida = [
    ("Divergências", formatar(divergencias.reindex(
        divergencias["impacto_rs"].abs().sort_values(ascending=False).index))),
    ("Impacto por cliente", resumo_cliente),
    ("Por data e filial", resumo_data),
    ("Por filial", resumo_filial),
    ("Por causa", resumo_causa),
    ("Por região", resumo_regiao),
    ("Sem preço na tabela", formatar(sem_preco)),
    ("Não auditado", formatar(nao_auditado)),
    ("Base completa", formatar(res)),
]

b1, b2, b3 = st.columns(3)
with b1:
    try:
        st.download_button(
            "📊 Baixar Excel completo", gerar_excel(abas_saida),
            file_name=f"auditoria_precos_{date.today():%Y%m%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True, type="primary",
        )
    except Exception as erro:  # noqa: BLE001
        st.warning(f"Excel indisponível neste ambiente ({erro}). "
                   "Baixe os CSVs ao lado.")
        st.download_button(
            "🗂️ Baixar CSVs (ZIP)", gerar_csvs_zip(abas_saida),
            file_name=f"auditoria_precos_{date.today():%Y%m%d}.zip",
            mime="application/zip", use_container_width=True,
        )
with b2:
    st.download_button("🖹 Baixar e-mail em HTML", html_email.encode("utf-8"),
                       file_name=f"email_auditoria_{date.today():%Y%m%d}.html",
                       mime="text/html", use_container_width=True)
with b3:
    st.download_button("📝 Baixar e-mail em texto", texto_email.encode("utf-8"),
                       file_name=f"email_auditoria_{date.today():%Y%m%d}.txt",
                       mime="text/plain", use_container_width=True)
